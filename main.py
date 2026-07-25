# -*- coding: utf-8 -*-
"""
Hisobchi — To'liq backend va Telegram bot (bitta faylda)
- Flask web-server
- Telegram bot (polling) alohida threadda
- JSON fayllarda ma'lumot saqlash (PostgreSQL o'rniga)
"""
import os
import hmac
import hashlib
import json
import io
import asyncio
import threading
import time
import logging
from datetime import datetime
from urllib.parse import parse_qsl
from pathlib import Path

import requests
from flask import Flask, request, jsonify, render_template, send_file

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

from aiogram import Bot, Dispatcher, Router, F
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import CommandStart
from aiogram.types import (
    Message,
    CallbackQuery,
    ReplyKeyboardMarkup,
    KeyboardButton,
    ReplyKeyboardRemove,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    WebAppInfo,
    MenuButtonWebApp,
)

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("hisobchi")


# ---------------------------------------------------------------------------
# Environment variables
# ---------------------------------------------------------------------------
BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
BOT_USERNAME = os.environ.get("BOT_USERNAME", "")
WEBAPP_URL = os.environ.get("WEBAPP_URL", "")

try:
    ADMIN_ID = int(os.environ.get("ADMIN_ID", "0"))
except ValueError:
    ADMIN_ID = 0


# ---------------------------------------------------------------------------
# JSON ma'lumotlar bazasi
# ---------------------------------------------------------------------------
DATA_DIR = Path("data")
DATA_DIR.mkdir(exist_ok=True)

USERS_FILE = DATA_DIR / "users.json"
TRANSACTIONS_FILE = DATA_DIR / "transactions.json"
DEBTS_FILE = DATA_DIR / "debts.json"

# Har bir fayl uchun ID counter
COUNTERS_FILE = DATA_DIR / "counters.json"


def _load_json(filepath, default=None):
    """JSON faylni yuklaydi, agar mavjud bo'lmasa default qaytaradi."""
    if default is None:
        default = {}
    try:
        if filepath.exists():
            with open(filepath, "r", encoding="utf-8") as f:
                return json.load(f)
        return default
    except (json.JSONDecodeError, OSError):
        return default


def _save_json(filepath, data):
    """Ma'lumotni JSON faylga saqlaydi."""
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _get_next_id(counter_name):
    """Berilgan counter uchun keyingi ID ni qaytaradi."""
    counters = _load_json(COUNTERS_FILE, {})
    next_id = counters.get(counter_name, 1)
    counters[counter_name] = next_id + 1
    _save_json(COUNTERS_FILE, counters)
    return next_id


def _to_iso(dt):
    """DateTime ni ISO formatga o'tkazadi."""
    if dt is None:
        return None
    return dt.isoformat() + "Z"


def _from_iso(iso_str):
    """ISO string dan datetime yaratadi."""
    if iso_str is None:
        return None
    return datetime.fromisoformat(iso_str.replace("Z", "+00:00"))


# ---------------------------------------------------------------------------
# Users
# ---------------------------------------------------------------------------
def get_user(telegram_id):
    users = _load_json(USERS_FILE, {})
    return users.get(str(telegram_id))


def upsert_user(telegram_id, phone, first_name, last_name, username):
    users = _load_json(USERS_FILE, {})
    users[str(telegram_id)] = {
        "telegram_id": telegram_id,
        "phone": phone,
        "first_name": first_name or "",
        "last_name": last_name or "",
        "username": username or "",
        "created_at": _to_iso(datetime.utcnow())
    }
    _save_json(USERS_FILE, users)
    return users[str(telegram_id)]


def get_all_users():
    users = _load_json(USERS_FILE, {})
    return list(users.values())


# ---------------------------------------------------------------------------
# Transactions
# ---------------------------------------------------------------------------
def get_transactions(telegram_id):
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    user_tx = all_tx.get(str(telegram_id), [])
    # created_at bo'yicha tartiblash
    return sorted(user_tx, key=lambda x: x.get("created_at", ""), reverse=True)


def add_transaction(telegram_id, type_, amount, category, note):
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    user_tx = all_tx.get(str(telegram_id), [])
    
    tx = {
        "id": _get_next_id("transaction"),
        "telegram_id": telegram_id,
        "type": type_,
        "amount": float(amount),
        "category": category or "",
        "note": note or "",
        "created_at": _to_iso(datetime.utcnow())
    }
    user_tx.append(tx)
    all_tx[str(telegram_id)] = user_tx
    _save_json(TRANSACTIONS_FILE, all_tx)
    return tx


def get_transaction(telegram_id, tx_id):
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    user_tx = all_tx.get(str(telegram_id), [])
    for tx in user_tx:
        if tx["id"] == tx_id:
            return tx
    return None


def delete_transaction(telegram_id, tx_id):
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    user_tx = all_tx.get(str(telegram_id), [])
    new_user_tx = [tx for tx in user_tx if tx["id"] != tx_id]
    if len(new_user_tx) != len(user_tx):
        all_tx[str(telegram_id)] = new_user_tx
        _save_json(TRANSACTIONS_FILE, all_tx)
        return True
    return False


def update_transaction(telegram_id, tx_id, type_, amount, category, note):
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    user_tx = all_tx.get(str(telegram_id), [])
    for tx in user_tx:
        if tx["id"] == tx_id:
            tx["type"] = type_
            tx["amount"] = float(amount)
            tx["category"] = category or ""
            tx["note"] = note or ""
            all_tx[str(telegram_id)] = user_tx
            _save_json(TRANSACTIONS_FILE, all_tx)
            return tx
    return None


def delete_all_transactions(telegram_id):
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    if str(telegram_id) in all_tx:
        del all_tx[str(telegram_id)]
        _save_json(TRANSACTIONS_FILE, all_tx)
        return True
    return False


# ---------------------------------------------------------------------------
# Debts
# ---------------------------------------------------------------------------
def get_debts(telegram_id):
    all_debts = _load_json(DEBTS_FILE, {})
    user_debts = all_debts.get(str(telegram_id), [])
    return sorted(user_debts, key=lambda x: (x.get("is_paid", False), x.get("created_at", "")), reverse=True)


def add_debt(telegram_id, direction, person_name, amount, note, is_payment=False):
    all_debts = _load_json(DEBTS_FILE, {})
    user_debts = all_debts.get(str(telegram_id), [])
    
    debt = {
        "id": _get_next_id("debt"),
        "telegram_id": telegram_id,
        "direction": direction,
        "person_name": person_name,
        "amount": float(amount),
        "note": note or "",
        "is_paid": False,
        "is_payment": bool(is_payment),
        "created_at": _to_iso(datetime.utcnow()),
        "paid_at": None
    }
    user_debts.append(debt)
    all_debts[str(telegram_id)] = user_debts
    _save_json(DEBTS_FILE, all_debts)
    return debt


def get_debt(telegram_id, debt_id):
    all_debts = _load_json(DEBTS_FILE, {})
    user_debts = all_debts.get(str(telegram_id), [])
    for debt in user_debts:
        if debt["id"] == debt_id:
            return debt
    return None


def update_debt(telegram_id, debt_id, person_name, amount, note):
    all_debts = _load_json(DEBTS_FILE, {})
    user_debts = all_debts.get(str(telegram_id), [])
    for debt in user_debts:
        if debt["id"] == debt_id:
            debt["person_name"] = person_name
            debt["amount"] = float(amount)
            debt["note"] = note or ""
            all_debts[str(telegram_id)] = user_debts
            _save_json(DEBTS_FILE, all_debts)
            return debt
    return None


def set_debt_paid(telegram_id, debt_id, is_paid):
    all_debts = _load_json(DEBTS_FILE, {})
    user_debts = all_debts.get(str(telegram_id), [])
    for debt in user_debts:
        if debt["id"] == debt_id:
            debt["is_paid"] = bool(is_paid)
            debt["paid_at"] = _to_iso(datetime.utcnow()) if is_paid else None
            all_debts[str(telegram_id)] = user_debts
            _save_json(DEBTS_FILE, all_debts)
            return debt
    return None


def delete_debt(telegram_id, debt_id):
    all_debts = _load_json(DEBTS_FILE, {})
    user_debts = all_debts.get(str(telegram_id), [])
    new_user_debts = [d for d in user_debts if d["id"] != debt_id]
    if len(new_user_debts) != len(user_debts):
        all_debts[str(telegram_id)] = new_user_debts
        _save_json(DEBTS_FILE, all_debts)
        return True
    return False


# ---------------------------------------------------------------------------
# Admin stats
# ---------------------------------------------------------------------------
def get_admin_stats():
    users = _load_json(USERS_FILE, {})
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    all_debts = _load_json(DEBTS_FILE, {})
    
    total_users = len(users)
    total_transactions = sum(len(tx_list) for tx_list in all_tx.values())
    total_debts = sum(len(d_list) for d_list in all_debts.values())
    
    # Shu oy uchun kirim/xarajat
    current_month = datetime.utcnow().strftime("%Y-%m")
    month_income = 0.0
    month_expense = 0.0
    
    for tx_list in all_tx.values():
        for tx in tx_list:
            created = tx.get("created_at", "")
            if created and created.startswith(current_month):
                if tx["type"] == "income":
                    month_income += float(tx["amount"])
                else:
                    month_expense += float(tx["amount"])
    
    # So'nggi 5 foydalanuvchi
    recent_users = sorted(
        users.values(),
        key=lambda x: x.get("created_at", ""),
        reverse=True
    )[:5]
    
    return {
        "total_users": total_users,
        "total_transactions": total_transactions,
        "total_debts": total_debts,
        "month_income": month_income,
        "month_expense": month_expense,
        "recent_users": recent_users,
    }


# ---------------------------------------------------------------------------
# Monthly summary
# ---------------------------------------------------------------------------
def get_monthly_summary(telegram_id):
    transactions = get_transactions(telegram_id)
    income = 0.0
    expense = 0.0
    categories = {}
    
    current_month = datetime.utcnow().strftime("%Y-%m")
    
    for t in transactions:
        created = t.get("created_at", "")
        if created and created.startswith(current_month):
            amount = float(t["amount"])
            if t["type"] == "income":
                income += amount
            else:
                expense += amount
                cat = t.get("category", "boshqa")
                categories[cat] = categories.get(cat, 0) + amount
    
    sorted_categories = sorted(categories.items(), key=lambda x: x[1], reverse=True)
    
    return {
        "income": income,
        "expense": expense,
        "balance": income - expense,
        "categories": sorted_categories,
    }


# ---------------------------------------------------------------------------
# Flask app
# ---------------------------------------------------------------------------
app = Flask(__name__)


# ---------------------------------------------------------------------------
# Telegram WebApp initData tekshiruvi
# ---------------------------------------------------------------------------
def verify_init_data(init_data: str):
    if not init_data or not BOT_TOKEN:
        return None
    try:
        pairs = dict(parse_qsl(init_data, strict_parsing=True))
    except ValueError:
        return None

    received_hash = pairs.pop("hash", None)
    if not received_hash:
        return None

    data_check_string = "\n".join(
        f"{k}={v}" for k, v in sorted(pairs.items())
    )
    secret_key = hmac.new(
        b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256
    ).digest()
    computed_hash = hmac.new(
        secret_key, data_check_string.encode(), hashlib.sha256
    ).hexdigest()

    if not hmac.compare_digest(computed_hash, received_hash):
        return None

    user_raw = pairs.get("user")
    if not user_raw:
        return None
    try:
        user = json.loads(user_raw)
    except json.JSONDecodeError:
        return None

    return user


def require_auth():
    init_data = request.headers.get("X-Telegram-Init-Data", "")
    user = verify_init_data(init_data)
    if user is None:
        return None, (jsonify({"error": "Telegram orqali kiring"}), 401)
    return user, None


def require_registered():
    user, err = require_auth()
    if err:
        return None, err
    db_user = get_user(user["id"])
    if not db_user:
        return None, (jsonify({
            "error": "not_registered",
            "message": "Avval botda ro'yxatdan o'ting"
        }), 403)
    return user, None


def is_admin_user(telegram_id):
    return ADMIN_ID != 0 and telegram_id == ADMIN_ID


def require_admin():
    user, err = require_registered()
    if err:
        return None, err
    if not is_admin_user(user["id"]):
        return None, (jsonify({
            "error": "faqat admin tahrirlashi/o'chirishi mumkin"
        }), 403)
    return user, None


# ---------------------------------------------------------------------------
# Sahifalar va API
# ---------------------------------------------------------------------------
def to_utc_iso(dt_str):
    """JSON'dan o'qilgan datetime stringni qaytaradi."""
    return dt_str


@app.route("/")
def index():
    return render_template("index.html", bot_username=BOT_USERNAME)


@app.route("/api/me")
def api_me():
    user, err = require_auth()
    if err:
        return err
    db_user = get_user(user["id"])
    if not db_user:
        return jsonify({"registered": False}), 200
    return jsonify({
        "registered": True,
        "first_name": db_user.get("first_name") or user.get("first_name", ""),
        "phone": db_user.get("phone", ""),
        "is_admin": is_admin_user(user["id"]),
    }), 200


@app.route("/api/transactions", methods=["GET"])
def api_list_transactions():
    user, err = require_registered()
    if err:
        return err
    rows = get_transactions(user["id"])
    result = [{
        "id": r["id"],
        "type": r["type"],
        "amount": float(r["amount"]),
        "category": r.get("category", ""),
        "note": r.get("note", ""),
        "created_at": r.get("created_at", ""),
    } for r in rows]
    return jsonify(result)


@app.route("/api/transactions", methods=["POST"])
def api_add_transaction():
    user, err = require_registered()
    if err:
        return err

    data = request.get_json(silent=True) or {}
    type_ = data.get("type")
    amount = data.get("amount")
    category = (data.get("category") or "").strip()[:100]
    note = (data.get("note") or "").strip()[:200]

    if type_ not in ("income", "expense"):
        return jsonify({"error": "noto'g'ri turi"}), 400
    try:
        amount = float(amount)
    except (TypeError, ValueError):
        return jsonify({"error": "noto'g'ri summa"}), 400
    if amount <= 0:
        return jsonify({"error": "summa 0 dan katta bo'lishi kerak"}), 400
    if type_ == "expense" and not category:
        return jsonify({"error": "kategoriya kerak"}), 400

    row = add_transaction(user["id"], type_, amount, category or "kirim", note)
    return jsonify({
        "id": row["id"],
        "type": row["type"],
        "amount": float(row["amount"]),
        "category": row.get("category", ""),
        "note": row.get("note", ""),
        "created_at": row.get("created_at", ""),
    }), 201


@app.route("/api/transactions/<int:tx_id>", methods=["PUT"])
def api_update_transaction(tx_id):
    user, err = require_admin()
    if err:
        return err

    data = request.get_json(silent=True) or {}
    type_ = data.get("type")
    amount = data.get("amount")
    category = (data.get("category") or "").strip()[:100]
    note = (data.get("note") or "").strip()[:200]

    if type_ not in ("income", "expense"):
        return jsonify({"error": "noto'g'ri turi"}), 400
    try:
        amount = float(amount)
    except (TypeError, ValueError):
        return jsonify({"error": "noto'g'ri summa"}), 400
    if amount <= 0:
        return jsonify({"error": "summa 0 dan katta bo'lishi kerak"}), 400
    if type_ == "expense" and not category:
        return jsonify({"error": "kategoriya kerak"}), 400

    row = update_transaction(user["id"], tx_id, type_, amount, category or "kirim", note)
    if not row:
        return jsonify({"error": "topilmadi"}), 404

    return jsonify({
        "id": row["id"],
        "type": row["type"],
        "amount": float(row["amount"]),
        "category": row.get("category", ""),
        "note": row.get("note", ""),
        "created_at": row.get("created_at", ""),
    }), 200


@app.route("/api/transactions/<int:tx_id>", methods=["DELETE"])
def api_delete_transaction(tx_id):
    user, err = require_admin()
    if err:
        return err
    ok = delete_transaction(user["id"], tx_id)
    if not ok:
        return jsonify({"error": "topilmadi"}), 404
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# API — Qarz daftari
# ---------------------------------------------------------------------------
def serialize_debt(row):
    return {
        "id": row["id"],
        "direction": row["direction"],
        "person_name": row["person_name"],
        "amount": float(row["amount"]),
        "note": row.get("note", ""),
        "is_paid": row.get("is_paid", False),
        "is_payment": row.get("is_payment", False),
        "created_at": row.get("created_at", ""),
        "paid_at": row.get("paid_at"),
    }


@app.route("/api/debts", methods=["GET"])
def api_list_debts():
    user, err = require_registered()
    if err:
        return err
    rows = get_debts(user["id"])
    return jsonify([serialize_debt(r) for r in rows])


@app.route("/api/debts", methods=["POST"])
def api_add_debt():
    user, err = require_registered()
    if err:
        return err

    data = request.get_json(silent=True) or {}
    direction = data.get("direction")
    person_name = (data.get("person_name") or "").strip()[:100]
    amount = data.get("amount")
    note = (data.get("note") or "").strip()[:200]
    is_payment = bool(data.get("is_payment", False))

    if direction not in ("given", "taken"):
        return jsonify({"error": "noto'g'ri turi"}), 400
    if not person_name:
        return jsonify({"error": "ism kiritilmagan"}), 400
    try:
        amount = float(amount)
    except (TypeError, ValueError):
        return jsonify({"error": "noto'g'ri summa"}), 400
    if amount <= 0:
        return jsonify({"error": "summa 0 dan katta bo'lishi kerak"}), 400

    row = add_debt(user["id"], direction, person_name, amount, note, is_payment)
    return jsonify(serialize_debt(row)), 201


@app.route("/api/debts/<int:debt_id>", methods=["PUT"])
def api_update_debt(debt_id):
    user, err = require_admin()
    if err:
        return err

    existing = get_debt(user["id"], debt_id)
    if not existing:
        return jsonify({"error": "topilmadi"}), 404

    data = request.get_json(silent=True) or {}

    if "is_paid" in data and len(data) == 1:
        row = set_debt_paid(user["id"], debt_id, bool(data["is_paid"]))
        return jsonify(serialize_debt(row))

    person_name = (data.get("person_name") or "").strip()[:100]
    amount = data.get("amount")
    note = (data.get("note") or "").strip()[:200]

    if not person_name:
        return jsonify({"error": "ism kiritilmagan"}), 400
    try:
        amount = float(amount)
    except (TypeError, ValueError):
        return jsonify({"error": "noto'g'ri summa"}), 400
    if amount <= 0:
        return jsonify({"error": "summa 0 dan katta bo'lishi kerak"}), 400

    row = update_debt(user["id"], debt_id, person_name, amount, note)
    if not row:
        return jsonify({"error": "topilmadi"}), 404
    return jsonify(serialize_debt(row))


@app.route("/api/debts/<int:debt_id>", methods=["DELETE"])
def api_delete_debt(debt_id):
    user, err = require_admin()
    if err:
        return err
    ok = delete_debt(user["id"], debt_id)
    if not ok:
        return jsonify({"error": "topilmadi"}), 404
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# Excel / PDF hisobot generatsiyasi
# ---------------------------------------------------------------------------
UZ_MONTHS = [
    "yanvar", "fevral", "mart", "aprel", "may", "iyun",
    "iyul", "avgust", "sentyabr", "oktyabr", "noyabr", "dekabr",
]

BRAND_DARK = "0A0E27"
BRAND_GREEN = "00B894"
BRAND_RED = "FF4757"


def _monthly_breakdown(transactions):
    months = {}
    for t in transactions:
        created = t.get("created_at", "")
        if not created:
            continue
        try:
            d = datetime.fromisoformat(created.replace("Z", "+00:00"))
        except:
            continue
        key = (d.year, d.month)
        if key not in months:
            months[key] = {"year": d.year, "month": d.month, "income": 0.0, "expense": 0.0}
        if t["type"] == "income":
            months[key]["income"] += float(t["amount"])
        else:
            months[key]["expense"] += float(t["amount"])
    return sorted(months.values(), key=lambda m: (m["year"], m["month"]), reverse=True)


def generate_excel_report(telegram_id, display_name):
    transactions = get_transactions(telegram_id)
    debts = get_debts(telegram_id)

    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color=BRAND_DARK, end_color=BRAND_DARK, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(bold=True, size=14, color=BRAND_DARK)
    thin_border = Border(*(Side(style="thin", color="DDDDDD"),) * 4)

    def style_header_row(ws, row_num, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def autosize(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws1 = wb.active
    ws1.title = "Oylik xulosa"
    ws1["A1"] = "Hisobchi — Moliyaviy hisobot"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Foydalanuvchi: {display_name}"
    ws1["A3"] = f"Yaratildi: {datetime.utcnow().strftime('%d.%m.%Y %H:%M')} (UTC)"

    headers = ["Oy", "Kirim (so'm)", "Xarajat (so'm)", "Balans (so'm)"]
    ws1.append([])
    ws1.append(headers)
    style_header_row(ws1, 5, len(headers))

    for m in _monthly_breakdown(transactions):
        balance = m["income"] - m["expense"]
        label = f"{UZ_MONTHS[m['month'] - 1]} {m['year']}"
        ws1.append([label, m["income"], m["expense"], balance])

    for row in ws1.iter_rows(min_row=6, max_row=ws1.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = thin_border
            if cell.column > 1:
                cell.number_format = "#,##0"

    autosize(ws1, [22, 18, 18, 18])

    ws2 = wb.create_sheet("Tranzaksiyalar")
    headers2 = ["Sana", "Turi", "Summa (so'm)", "Kategoriya", "Izoh"]
    ws2.append(headers2)
    style_header_row(ws2, 1, len(headers2))

    for t in transactions:
        created = t.get("created_at", "")
        try:
            dt = datetime.fromisoformat(created.replace("Z", "+00:00"))
            date_str = dt.strftime("%d.%m.%Y %H:%M")
        except:
            date_str = created
        ws2.append([
            date_str,
            "Kirim" if t["type"] == "income" else "Xarajat",
            float(t["amount"]),
            t.get("category", ""),
            t.get("note", ""),
        ])

    for row in ws2.iter_rows(min_row=2, max_row=max(ws2.max_row, 2), min_col=1, max_col=5):
        for cell in row:
            cell.border = thin_border
            if cell.column == 3:
                cell.number_format = "#,##0"

    autosize(ws2, [18, 12, 16, 20, 30])

    ws3 = wb.create_sheet("Qarz daftari")
    headers3 = ["Ism", "Yo'nalish", "Summa (so'm)", "Holat", "Sana", "Izoh"]
    ws3.append(headers3)
    style_header_row(ws3, 1, len(headers3))

    for d in debts:
        created = d.get("created_at", "")
        try:
            dt = datetime.fromisoformat(created.replace("Z", "+00:00"))
            date_str = dt.strftime("%d.%m.%Y")
        except:
            date_str = created
        ws3.append([
            d["person_name"],
            "Menga qarzdor" if d["direction"] == "given" else "Men qarzdorman",
            float(d["amount"]),
            "To'landi" if d.get("is_paid") else "To'lanmagan",
            date_str,
            d.get("note", ""),
        ])

    for row in ws3.iter_rows(min_row=2, max_row=max(ws3.max_row, 2), min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
            if cell.column == 3:
                cell.number_format = "#,##0"

    autosize(ws3, [18, 16, 16, 14, 14, 30])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_pdf_report(telegram_id, display_name):
    transactions = get_transactions(telegram_id)
    debts = get_debts(telegram_id)
    months = _monthly_breakdown(transactions)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=18 * mm, bottomMargin=18 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleUz", parent=styles["Title"], fontSize=18,
        textColor=colors.HexColor("#0A0E27"), alignment=TA_CENTER, spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "MetaUz", parent=styles["Normal"], fontSize=9,
        textColor=colors.HexColor("#666666"), alignment=TA_CENTER, spaceAfter=14,
    )
    section_style = ParagraphStyle(
        "SectionUz", parent=styles["Heading2"], fontSize=13,
        textColor=colors.HexColor("#0A0E27"), spaceBefore=14, spaceAfter=8,
    )

    def fmt_num(n):
        return f"{n:,.0f}".replace(",", " ")

    elements = [
        Paragraph("Hisobchi — Moliyaviy hisobot", title_style),
        Paragraph(
            f"Foydalanuvchi: {display_name} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Yaratildi: {datetime.utcnow().strftime('%d.%m.%Y %H:%M')} (UTC)",
            meta_style,
        ),
    ]

    total_income = sum(float(t["amount"]) for t in transactions if t["type"] == "income")
    total_expense = sum(float(t["amount"]) for t in transactions if t["type"] == "expense")

    elements.append(Paragraph("Umumiy holat", section_style))
    summary_data = [
        ["Jami kirim", "Jami xarajat", "Balans"],
        [f"{fmt_num(total_income)} so'm", f"{fmt_num(total_expense)} so'm",
         f"{fmt_num(total_income - total_expense)} so'm"],
    ]
    summary_table = Table(summary_data, colWidths=[55 * mm, 55 * mm, 55 * mm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0E27")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
    ]))
    elements.append(summary_table)

    if months:
        elements.append(Paragraph("Oylik xulosa", section_style))
        month_rows = [["Oy", "Kirim", "Xarajat", "Balans"]]
        for m in months:
            balance = m["income"] - m["expense"]
            month_rows.append([
                f"{UZ_MONTHS[m['month'] - 1]} {m['year']}",
                fmt_num(m["income"]), fmt_num(m["expense"]), fmt_num(balance),
            ])
        month_table = Table(month_rows, colWidths=[45 * mm, 40 * mm, 40 * mm, 40 * mm])
        month_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0E27")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F6FA")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(month_table)

    if debts:
        elements.append(Paragraph("Qarz daftari", section_style))
        debt_rows = [["Ism", "Yo'nalish", "Summa", "Holat"]]
        for d in debts:
            debt_rows.append([
                d["person_name"],
                "Menga qarzdor" if d["direction"] == "given" else "Men qarzdorman",
                fmt_num(float(d["amount"])),
                "To'landi" if d.get("is_paid") else "To'lanmagan",
            ])
        debt_table = Table(debt_rows, colWidths=[45 * mm, 40 * mm, 35 * mm, 45 * mm])
        debt_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0E27")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 0), (2, -1), "RIGHT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F6FA")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(debt_table)

    if transactions:
        elements.append(Spacer(1, 6))
        elements.append(Paragraph("Barcha tranzaksiyalar", section_style))
        tx_rows = [["Sana", "Turi", "Summa", "Kategoriya"]]
        for t in transactions:
            created = t.get("created_at", "")
            try:
                dt = datetime.fromisoformat(created.replace("Z", "+00:00"))
                date_str = dt.strftime("%d.%m.%Y %H:%M")
            except:
                date_str = created
            tx_rows.append([
                date_str,
                "Kirim" if t["type"] == "income" else "Xarajat",
                fmt_num(float(t["amount"])),
                t.get("category", ""),
            ])
        tx_table = Table(tx_rows, colWidths=[38 * mm, 25 * mm, 32 * mm, 65 * mm], repeatRows=1)
        tx_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0E27")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 0), (2, -1), "RIGHT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F6FA")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(tx_table)

    doc.build(elements)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# API — Excel / PDF hisobot
# ---------------------------------------------------------------------------
def send_telegram_document(chat_id, filename, file_bytes, caption=""):
    url = f"https://api.telegram.org/bot{BOT_TOKEN}/sendDocument"
    files = {"document": (filename, file_bytes)}
    data = {"chat_id": chat_id, "caption": caption}
    resp = requests.post(url, data=data, files=files, timeout=30)
    resp.raise_for_status()
    result = resp.json()
    if not result.get("ok"):
        raise RuntimeError(result.get("description", "Telegram xatosi"))
    return result


@app.route("/api/export/send", methods=["POST"])
def api_export_send():
    user, err = require_registered()
    if err:
        return err

    data = request.get_json(silent=True) or {}
    fmt = data.get("format")
    if fmt not in ("excel", "pdf"):
        return jsonify({"error": "noto'g'ri format"}), 400

    db_user = get_user(user["id"])
    display_name = db_user.get("first_name") or user.get("first_name", "Foydalanuvchi")

    try:
        if fmt == "excel":
            buf = generate_excel_report(user["id"], display_name)
            filename = f"hisobchi_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
        else:
            buf = generate_pdf_report(user["id"], display_name)
            filename = f"hisobchi_{datetime.utcnow().strftime('%Y%m%d')}.pdf"

        send_telegram_document(
            user["id"], filename, buf.read(),
            caption="📊 Moliyaviy hisobotingiz tayyor."
        )
    except Exception as e:
        return jsonify({"error": f"hisobotni yuborib bo'lmadi: {e}"}), 500

    return jsonify({"success": True})


@app.route("/api/export/excel")
def api_export_excel():
    user, err = require_registered()
    if err:
        return err
    db_user = get_user(user["id"])
    display_name = db_user.get("first_name") or user.get("first_name", "Foydalanuvchi")

    buf = generate_excel_report(user["id"], display_name)
    filename = f"hisobchi_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@app.route("/api/export/pdf")
def api_export_pdf():
    user, err = require_registered()
    if err:
        return err
    db_user = get_user(user["id"])
    display_name = db_user.get("first_name") or user.get("first_name", "Foydalanuvchi")

    buf = generate_pdf_report(user["id"], display_name)
    filename = f"hisobchi_{datetime.utcnow().strftime('%Y%m%d')}.pdf"
    return send_file(
        buf,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )


# ---------------------------------------------------------------------------
# Admin — foydalanuvchilarni boshqarish
# ---------------------------------------------------------------------------
def serialize_admin_user(u, all_tx, all_debts):
    uid = str(u["telegram_id"])
    return {
        "telegram_id": u["telegram_id"],
        "first_name": u.get("first_name", ""),
        "last_name": u.get("last_name", ""),
        "phone": u.get("phone", ""),
        "username": u.get("username", ""),
        "created_at": u.get("created_at", ""),
        "tx_count": len(all_tx.get(uid, [])),
        "debt_count": len(all_debts.get(uid, [])),
        "is_admin": is_admin_user(u["telegram_id"]),
    }


@app.route("/api/admin/users", methods=["GET"])
def api_admin_list_users():
    user, err = require_admin()
    if err:
        return err

    users = _load_json(USERS_FILE, {})
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    all_debts = _load_json(DEBTS_FILE, {})

    result = [
        serialize_admin_user(u, all_tx, all_debts)
        for u in sorted(users.values(), key=lambda x: x.get("created_at", ""), reverse=True)
    ]
    return jsonify(result)


@app.route("/api/admin/users/<int:user_id>", methods=["PUT"])
def api_admin_update_user(user_id):
    user, err = require_admin()
    if err:
        return err

    users = _load_json(USERS_FILE, {})
    if str(user_id) not in users:
        return jsonify({"error": "foydalanuvchi topilmadi"}), 404

    data = request.get_json(silent=True) or {}
    first_name = (data.get("first_name") or "").strip()[:100]
    last_name = (data.get("last_name") or "").strip()[:100]
    phone = (data.get("phone") or "").strip()[:30]

    if not first_name:
        return jsonify({"error": "ism kiritilmagan"}), 400

    users[str(user_id)]["first_name"] = first_name
    users[str(user_id)]["last_name"] = last_name
    users[str(user_id)]["phone"] = phone
    _save_json(USERS_FILE, users)

    all_tx = _load_json(TRANSACTIONS_FILE, {})
    all_debts = _load_json(DEBTS_FILE, {})
    return jsonify(serialize_admin_user(users[str(user_id)], all_tx, all_debts))


@app.route("/api/admin/users/<int:user_id>", methods=["DELETE"])
def api_admin_delete_user(user_id):
    user, err = require_admin()
    if err:
        return err
    
    users = _load_json(USERS_FILE, {})
    if str(user_id) not in users:
        return jsonify({"error": "foydalanuvchi topilmadi"}), 404
    
    # Foydalanuvchini o'chirish
    del users[str(user_id)]
    _save_json(USERS_FILE, users)
    
    # Tranzaksiyalarni o'chirish
    all_tx = _load_json(TRANSACTIONS_FILE, {})
    if str(user_id) in all_tx:
        del all_tx[str(user_id)]
        _save_json(TRANSACTIONS_FILE, all_tx)
    
    # Qarzlarni o'chirish
    all_debts = _load_json(DEBTS_FILE, {})
    if str(user_id) in all_debts:
        del all_debts[str(user_id)]
        _save_json(DEBTS_FILE, all_debts)
    
    return jsonify({"success": True})


# ---------------------------------------------------------------------------
# Telegram Bot
# ---------------------------------------------------------------------------
router = Router()
dp = Dispatcher()
dp.include_router(router)


def contact_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text="📱 Raqamni yuborish", request_contact=True)]],
        resize_keyboard=True,
        one_time_keyboard=True,
    )


def webapp_keyboard(is_admin: bool = False) -> InlineKeyboardMarkup:
    if is_admin:
        # Admin uchun faqat bitta tugma — to'g'ridan-to'g'ri Admin panelga olib boradi.
        return InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="🛠 Admin panelni ochish",
                    web_app=WebAppInfo(url=WEBAPP_URL),
                )
            ],
        ])

    rows = [
        [
            InlineKeyboardButton(
                text="📊 Hisobchini ochish",
                web_app=WebAppInfo(url=WEBAPP_URL),
            )
        ],
        [
            InlineKeyboardButton(
                text="📅 Oylik hisobot",
                callback_data="monthly_report",
            )
        ],
    ]
    return InlineKeyboardMarkup(inline_keyboard=rows)


def format_monthly_report(summary: dict) -> str:
    now = datetime.utcnow()
    month_name = UZ_MONTHS[now.month - 1]

    def fmt(n):
        return f"{n:,.0f}".replace(",", " ")

    lines = [f"📅 <b>{month_name.capitalize()} oyi uchun hisobot</b>\n"]
    lines.append(f"➕ Kirim: <b>{fmt(summary['income'])} so'm</b>")
    lines.append(f"➖ Xarajat: <b>{fmt(summary['expense'])} so'm</b>")
    lines.append(f"💰 Balans: <b>{fmt(summary['balance'])} so'm</b>")

    if summary["categories"]:
        lines.append("\n<b>Xarajatlar taqsimoti:</b>")
        total_expense = summary["expense"] or 1
        for cat, val in summary["categories"][:8]:
            pct = round(val / total_expense * 100)
            lines.append(f"• {cat} — {fmt(val)} so'm ({pct}%)")
    else:
        lines.append("\nBu oyda hali xarajat kiritilmagan.")

    return "\n".join(lines)


@router.message(CommandStart())
async def cmd_start(message: Message):
    is_admin = ADMIN_ID != 0 and message.from_user.id == ADMIN_ID
    user = get_user(message.from_user.id)

    if user:
        if is_admin:
            greeting = (
                f"Salom, <b>{message.from_user.first_name}</b>! 👋\n\n"
                "Siz <b>admin</b> sifatida kirdingiz. Quyidagi tugma orqali "
                "faqat foydalanuvchilarni boshqarish uchun Admin panelni oching."
            )
        else:
            greeting = (
                f"Salom, <b>{message.from_user.first_name}</b>! 👋\n\n"
                "Hisobchi tayyor — kirim va xarajatlaringizni boshqarish uchun "
                "quyidagi tugmani bosing."
            )
        await message.answer(greeting, reply_markup=webapp_keyboard(is_admin))
        return

    if is_admin:
        # Admin ro'yxatdan o'tmagan bo'lsa ham, telefon so'ramasdan
        # to'g'ridan-to'g'ri Admin panelga kira oladi.
        upsert_user(
            telegram_id=message.from_user.id,
            phone="",
            first_name=message.from_user.first_name,
            last_name=message.from_user.last_name,
            username=message.from_user.username,
        )
        await message.answer(
            f"Salom, <b>{message.from_user.first_name}</b>! 👋\n\n"
            "Siz <b>admin</b> sifatida kirdingiz. Quyidagi tugma orqali "
            "Admin panelni oching.",
            reply_markup=webapp_keyboard(True),
        )
        return

    await message.answer(
        "Assalomu alaykum! 👋\n\n"
        "<b>Hisobchi</b> botiga xush kelibsiz — bu bot orqali kirim va "
        "xarajatlaringizni qulay tarzda hisoblab borishingiz mumkin.\n\n"
        "Davom etish uchun avval telefon raqamingizni tasdiqlashingiz kerak. "
        "Pastdagi tugmani bosing 👇",
        reply_markup=contact_keyboard(),
    )


@router.message(F.contact)
async def on_contact(message: Message):
    contact = message.contact
    if contact.user_id != message.from_user.id:
        await message.answer(
            "❗️ Iltimos, faqat <b>o'zingizning</b> raqamingizni yuboring.",
            reply_markup=contact_keyboard(),
        )
        return

    upsert_user(
        telegram_id=message.from_user.id,
        phone=contact.phone_number,
        first_name=message.from_user.first_name,
        last_name=message.from_user.last_name,
        username=message.from_user.username,
    )

    is_admin = ADMIN_ID != 0 and message.from_user.id == ADMIN_ID

    await message.answer(
        "✅ Ro'yxatdan muvaffaqiyatli o'tdingiz!",
        reply_markup=ReplyKeyboardRemove(),
    )
    await message.answer(
        "Admin panelni ochish uchun tugmani bosing 👇" if is_admin
        else "Endi Hisobchidan foydalanishingiz mumkin 👇",
        reply_markup=webapp_keyboard(is_admin),
    )


@router.message()
async def block_unregistered(message: Message):
    user = get_user(message.from_user.id)
    is_admin = ADMIN_ID != 0 and message.from_user.id == ADMIN_ID
    if user:
        await message.answer(
            "Admin panelni ochish uchun tugmani bosing 👇" if is_admin
            else "Hisobchini ochish uchun tugmani bosing 👇",
            reply_markup=webapp_keyboard(is_admin),
        )
    else:
        await message.answer(
            "Davom etish uchun avval telefon raqamingizni yuboring 👇",
            reply_markup=contact_keyboard(),
        )


@router.callback_query(F.data == "monthly_report")
async def on_monthly_report(callback: CallbackQuery):
    user = get_user(callback.from_user.id)
    if not user:
        await callback.answer("Avval ro'yxatdan o'ting", show_alert=True)
        return

    summary = get_monthly_summary(callback.from_user.id)
    text = format_monthly_report(summary)
    await callback.message.answer(text)
    await callback.answer()


async def _run_polling():
    bot = Bot(token=BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    try:
        await bot.delete_webhook(drop_pending_updates=True)
        await bot.set_chat_menu_button(
            menu_button=MenuButtonWebApp(
                text="📊 Hisobchi",
                web_app=WebAppInfo(url=WEBAPP_URL),
            )
        )
        await dp.start_polling(bot, handle_signals=False)
    finally:
        await bot.session.close()


def start_bot():
    """Alohida thread ichidan chaqiriladi."""
    asyncio.run(_run_polling())


# ---------------------------------------------------------------------------
# Asosiy ishga tushirish
# ---------------------------------------------------------------------------
def run_bot_thread():
    while True:
        try:
            logger.info("Bot polling boshlanmoqda...")
            start_bot()
        except Exception:
            logger.exception("Bot xato bilan to'xtadi, 5 soniyadan keyin qayta urinamiz")
            time.sleep(5)


if __name__ == "__main__":
    # `data/` papkasini yaratamiz
    DATA_DIR.mkdir(exist_ok=True)
    logger.info("JSON ma'lumotlar bazasi tayyor")

    bot_thread = threading.Thread(target=run_bot_thread, daemon=True)
    bot_thread.start()
    logger.info("Bot thread started")

    port = int(os.environ.get("PORT", 5000))
    logger.info(f"Flask server starting on port {port}")
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True, use_reloader=False)
